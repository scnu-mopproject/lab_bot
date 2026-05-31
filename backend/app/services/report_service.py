"""报表导出：按条件查询并生成 Excel(xlsx)。"""
import io
from datetime import date, datetime, time, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.booking import Booking
from app.models.repair import Repair
from app.models.room import Room
from app.models.user import User

BOOKING_STATUS_CN = {
    "pending": "待审批", "approved": "已通过", "rejected": "已驳回", "cancelled": "已取消",
}
REPAIR_STATUS_CN = {
    "submitted": "待受理", "assigned": "已分配", "in_progress": "维修中", "done": "已完成", "closed": "已关闭",
}
SOURCE_CN = {"user": "师生预约", "course": "课表占用"}

_HEADER_FILL = PatternFill("solid", fgColor="2B6CB0")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _day_bounds(start: date | None, end: date | None) -> tuple[datetime | None, datetime | None]:
    s = datetime.combine(start, time.min) if start else None
    e = datetime.combine(end, time.max) if end else None
    return s, e


def _autofit(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _fmt(dt) -> str:
    return dt.strftime("%Y-%m-%d %H:%M") if isinstance(dt, datetime) else (dt or "")


def build_bookings_report(db: Session, *, start: date | None = None, end: date | None = None,
                          status: str | None = None, room_id: int | None = None,
                          source: str | None = None) -> bytes:
    s, e = _day_bounds(start, end)
    stmt = select(Booking).order_by(Booking.start_time)
    if s:
        stmt = stmt.where(Booking.start_time >= s)
    if e:
        stmt = stmt.where(Booking.start_time <= e)
    if status:
        stmt = stmt.where(Booking.status == status)
    if room_id:
        stmt = stmt.where(Booking.room_id == room_id)
    if source:
        stmt = stmt.where(Booking.source == source)
    rows = db.scalars(stmt).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "预约记录"
    _write_header(ws, ["预约号", "实验室", "预约人", "开始时间", "结束时间", "用途", "来源", "状态", "提交时间"])
    for b in rows:
        room = db.get(Room, b.room_id)
        u = db.get(User, b.user_id)
        ws.append([
            b.id, room.name if room else "", u.name if u else "",
            _fmt(b.start_time), _fmt(b.end_time), b.purpose or b.course_name or "",
            SOURCE_CN.get(b.source, b.source), BOOKING_STATUS_CN.get(b.status, b.status),
            _fmt(b.created_at),
        ])
    _autofit(ws, [8, 22, 12, 18, 18, 24, 12, 10, 18])
    return _to_bytes(wb)


def build_repairs_report(db: Session, *, start: date | None = None, end: date | None = None,
                         status: str | None = None, room_id: int | None = None) -> bytes:
    s, e = _day_bounds(start, end)
    stmt = select(Repair).order_by(Repair.created_at)
    if s:
        stmt = stmt.where(Repair.created_at >= s)
    if e:
        stmt = stmt.where(Repair.created_at <= e)
    if status:
        stmt = stmt.where(Repair.status == status)
    if room_id:
        stmt = stmt.where(Repair.room_id == room_id)
    rows = db.scalars(stmt).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "报修记录"
    _write_header(ws, ["工单号", "实验室", "设备", "故障描述", "报修人", "处理人", "状态", "提交时间", "更新时间"])
    for r in rows:
        room = db.get(Room, r.room_id)
        reporter = db.get(User, r.reporter_id)
        handler = db.get(User, r.handler_id) if r.handler_id else None
        ws.append([
            r.id, room.name if room else "", r.device_name, r.description,
            reporter.name if reporter else "", handler.name if handler else "",
            REPAIR_STATUS_CN.get(r.status, r.status), _fmt(r.created_at), _fmt(r.updated_at),
        ])
    _autofit(ws, [8, 22, 16, 36, 12, 12, 10, 18, 18])
    return _to_bytes(wb)


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
