"""知识文档：解析 -> 切片 -> 向量化 -> 入库；zip 批量；去重；分页。

支持 txt / md / pdf / docx / xlsx；压缩包仅支持 zip。
"""
import hashlib
import io
import re
import zipfile

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.models.document import Document, DocumentChunk
from app.services.ai.embedding import get_embedding_provider

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50
SUPPORTED_EXT = (".txt", ".md", ".markdown", ".pdf", ".docx", ".xlsx")

# zip 解压防护
MAX_ZIP_FILES = 100
MAX_TOTAL_UNCOMPRESSED = 200 * 1024 * 1024  # 200MB
MAX_SINGLE_FILE = 20 * 1024 * 1024          # 20MB

_SENT_SPLIT = re.compile(r"(?<=[。！？!?；;\n])")


def parse_file(content: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith((".txt", ".md", ".markdown")):
        return content.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(content))
        parts = [p.text for p in d.paragraphs]
        for tbl in d.tables:
            for row in tbl.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" ".join(cells))
        return "\n".join(lines)
    raise ValueError("仅支持 txt / md / pdf / docx / xlsx 文件")


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        return []
    sentences = [s.strip() for s in _SENT_SPLIT.split(text) if s.strip()]
    chunks: list[str] = []
    cur = ""
    for s in sentences:
        if len(s) > size:
            if cur:
                chunks.append(cur)
                cur = ""
            for i in range(0, len(s), size - overlap):
                chunks.append(s[i:i + size])
            continue
        if len(cur) + len(s) <= size:
            cur += s
        else:
            chunks.append(cur)
            cur = (cur[-overlap:] if overlap else "") + s
    if cur:
        chunks.append(cur)
    return chunks


def _hash_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _index_one(db: Session, *, content: bytes, filename: str, title: str | None,
               uploaded_by: int | None) -> tuple[str, Document | None]:
    """解析并入库单个文件（不 commit）。返回 (status, doc)，status: created|duplicate|empty。"""
    text = parse_file(content, filename)
    chunks = chunk_text(text)
    if not chunks:
        return "empty", None
    content_hash = _hash_text(text)
    existing = db.scalar(select(Document).where(Document.content_hash == content_hash))
    if existing:
        return "duplicate", existing

    provider = get_embedding_provider()
    vectors = provider.embed(chunks)
    doc = Document(
        title=title or filename, filename=filename, uploaded_by=uploaded_by,
        content_hash=content_hash, chunk_count=len(chunks), embedding_model=provider.name,
    )
    db.add(doc)
    db.flush()
    for i, (chunk, vec) in enumerate(zip(chunks, vectors)):
        db.add(DocumentChunk(document_id=doc.id, chunk_index=i, content=chunk,
                             embedding=vec, embedding_model=provider.name))
    return "created", doc


def index_single(db: Session, *, content: bytes, filename: str, title: str | None,
                 uploaded_by: int | None) -> tuple[str, Document | None]:
    status, doc = _index_one(db, content=content, filename=filename, title=title,
                             uploaded_by=uploaded_by)
    if status == "empty":
        raise ValueError("文档解析后无有效文本内容")
    if status == "created":
        db.commit()
        db.refresh(doc)
    return status, doc


def _zip_entry_name(info: zipfile.ZipInfo) -> str:
    """处理 zip 中文文件名编码：UTF-8 标志位未设时按 GBK 兜底。"""
    name = info.filename
    if info.flag_bits & 0x800:
        return name
    try:
        return name.encode("cp437").decode("gbk")
    except Exception:  # noqa: BLE001
        return name


def index_zip(db: Session, *, content: bytes, uploaded_by: int | None) -> dict:
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise ValueError("不是有效的 zip 压缩包")

    created = duplicated = 0
    skipped: list[str] = []
    docs: list[Document] = []
    total_size = 0
    count = 0

    for info in zf.infolist():
        if info.is_dir():
            continue
        name = _zip_entry_name(info)
        base = name.split("/")[-1]
        if not base or base.startswith(".") or "__MACOSX" in name:
            continue
        if not name.lower().endswith(SUPPORTED_EXT):
            skipped.append(f"{base}：不支持的格式")
            continue
        if info.file_size > MAX_SINGLE_FILE:
            skipped.append(f"{base}：单文件超过 20MB")
            continue
        count += 1
        total_size += info.file_size
        if count > MAX_ZIP_FILES:
            skipped.append(f"超过最多 {MAX_ZIP_FILES} 个文件，其余未处理")
            break
        if total_size > MAX_TOTAL_UNCOMPRESSED:
            skipped.append("累计解压体积超过 200MB，其余未处理")
            break

        try:
            data = zf.read(info)
            status, doc = _index_one(db, content=data, filename=base, title=base,
                                     uploaded_by=uploaded_by)
        except Exception:  # noqa: BLE001
            skipped.append(f"{base}：解析失败")
            continue
        if status == "created":
            created += 1
            docs.append(doc)
        elif status == "duplicate":
            duplicated += 1
        else:
            skipped.append(f"{base}：无有效内容")

    db.commit()
    for d in docs:
        db.refresh(d)
    return {"created": created, "duplicated": duplicated, "skipped": skipped, "documents": docs}


def delete_document(db: Session, doc_id: int) -> bool:
    doc = db.get(Document, doc_id)
    if not doc:
        return False
    db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc_id))
    db.delete(doc)
    db.commit()
    return True


def reindex_all(db: Session) -> dict:
    """用当前向量化方案重新计算所有切片的向量（切换 embedding 后调用，无需重新上传）。"""
    provider = get_embedding_provider()
    chunks = list(db.scalars(select(DocumentChunk)).all())
    docs = list(db.scalars(select(Document)).all())
    if chunks:
        vectors = provider.embed([c.content for c in chunks])
        for c, v in zip(chunks, vectors):
            c.embedding = v
            c.embedding_model = provider.name
    for d in docs:
        d.embedding_model = provider.name
    db.commit()
    return {"documents": len(docs), "chunks": len(chunks), "embedding_model": provider.name}


def list_documents(db: Session, *, page: int = 1, page_size: int = 20,
                   keyword: str | None = None) -> tuple[list[Document], int]:
    stmt = select(Document)
    if keyword:
        stmt = stmt.where(Document.title.like(f"%{keyword}%"))
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    items = list(db.scalars(
        stmt.order_by(Document.created_at.desc())
        .offset((page - 1) * page_size).limit(page_size)
    ).all())
    return items, total
